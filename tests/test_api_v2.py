import asyncio
import json
from concurrent.futures import Future
from concurrent.futures import TimeoutError as FutureTimeout

import pytest
from fastapi.testclient import TestClient
from pebble import ProcessExpired, ProcessPool

import api_v2

client = TestClient(api_v2.app)


class InlinePool:
    """Runs the worker function in-process; optionally fails chosen filenames."""

    def __init__(self, fail=None):
        self.fail = fail or {}
        self.scheduled = []

    def schedule(self, fn, args=(), timeout=None):
        content, name, skills = args
        self.scheduled.append(name)
        f = Future()
        if name in self.fail:
            f.set_exception(self.fail[name])
        else:
            f.set_result(fn(*args))
        return f


@pytest.fixture(autouse=True)
def fake_pool():
    api_v2.app.state.pool = InlinePool()
    api_v2._active_requests = 0
    yield api_v2.app.state.pool


def post(files, skills="python,sql", headers=None, path="/rank-resumes"):
    return client.post(path, data={"skills": skills}, files=files, headers=headers or {})


def lines(resp):
    return [json.loads(l) for l in resp.text.splitlines() if l]


def test_health_both_paths():
    assert client.get("/health").json()["version"] == 2
    assert client.get("/v2/health").status_code == 200


def test_valid_upload_on_both_routes(make_docx):
    for path in ("/rank-resumes", "/v2/rank-resumes"):
        r = post([("files", ("a.docx", make_docx(["Python and SQL"])))], path=path)
        (res,) = lines(r)
        assert res["status"] == "success" and res["score"] == 2


def test_path_style_filenames_preserved(make_docx):
    data = make_docx(["Python"])
    r = post([("files", ("a/cv.docx", data)), ("files", ("b/cv.docx", data))])
    assert sorted(x["filename"] for x in lines(r)) == ["a/cv.docx", "b/cv.docx"]


def test_docx_header_and_footer_are_searched():
    import io
    import docx
    d = docx.Document()
    d.add_paragraph("body")
    d.sections[0].header.paragraphs[0].text = "Python"
    d.sections[0].footer.paragraphs[0].text = "SQL"
    buf = io.BytesIO()
    d.save(buf)
    r = post([("files", ("h.docx", buf.getvalue()))])
    assert lines(r)[0]["matched_skills"] == ["python", "sql"]


# --- validation ---
def test_skill_limits(make_docx):
    f = [("files", ("a.docx", make_docx(["x"])))]
    assert post(f, skills=" , ").status_code == 400
    assert post(f, skills=",".join(f"s{i}" for i in range(api_v2.MAX_SKILLS + 1))).status_code == 400
    assert post(f, skills="x" * (api_v2.MAX_SKILL_LENGTH + 1)).status_code == 400


def test_too_many_files(make_docx):
    data = make_docx(["x"])
    files = [("files", (f"{i}.docx", data)) for i in range(api_v2.MAX_RESUME_COUNT + 1)]
    assert post(files).status_code == 400


def test_oversized_file(monkeypatch, make_docx):
    monkeypatch.setattr(api_v2, "MAX_FILE_SIZE", 10)
    r = post([("files", ("a.docx", make_docx(["Python"])))])
    assert r.status_code == 413 or lines(r)[0]["status"] == "error"


def test_request_size_cap(monkeypatch, make_docx):
    monkeypatch.setattr(api_v2, "MAX_REQUEST_BYTES", 100)
    r = post([("files", ("a.docx", make_docx(["Python"])))])
    assert r.status_code == 413


def test_content_must_match_extension():
    r = post([("files", ("fake.pdf", b"this is not a pdf"))])
    (res,) = lines(r)
    assert res["status"] == "error" and "does not match" in res["error"]


# --- auth and concurrency ---
def test_api_key(monkeypatch, make_docx):
    monkeypatch.setenv("API_KEY", "secret")
    f = lambda: [("files", ("a.docx", make_docx(["Python"])))]
    assert post(f()).status_code == 401
    assert post(f(), headers={"X-API-Key": "wrong"}).status_code == 401
    assert post(f(), headers={"X-API-Key": "secret"}).status_code == 200


def test_no_api_key_configured_means_open(make_docx):
    assert post([("files", ("a.docx", make_docx(["Python"])))]).status_code == 200


def test_busy_server_returns_429(make_docx):
    api_v2._active_requests = api_v2.MAX_CONCURRENT_REQUESTS
    r = post([("files", ("a.docx", make_docx(["Python"])))])
    assert r.status_code == 429 and r.headers["retry-after"]


def test_active_counter_released(make_docx):
    post([("files", ("a.docx", make_docx(["Python"])))]).text
    assert api_v2._active_requests == 0


# --- failure isolation ---
def test_timeout_crash_and_unexpected_errors_are_per_file(fake_pool, make_docx):
    fake_pool.fail = {
        "slow.docx": FutureTimeout(),
        "crash.docx": ProcessExpired("died"),
        "boom.docx": RuntimeError("secret internals"),
    }
    data = make_docx(["Python"])
    r = post([("files", (n, data)) for n in ("ok.docx", "slow.docx", "crash.docx", "boom.docx")])
    by = {x["filename"]: x for x in lines(r)}
    assert by["ok.docx"]["status"] == "success"
    assert "timed out" in by["slow.docx"]["error"]
    assert "crashed" in by["crash.docx"]["error"]
    assert by["boom.docx"]["error"] == "Internal error."
    assert "secret" not in r.text


def test_cancellation_cancels_the_pool_future():
    fut = Future()  # never completes

    class HangingPool:
        def schedule(self, *a, **k):
            return fut

    async def scenario():
        task = asyncio.ensure_future(api_v2.run_one(HangingPool(), b"x", "a.docx", ["python"]))
        await asyncio.sleep(0.05)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task

    asyncio.run(scenario())
    assert fut.cancelled()


def test_bounded_window_processes_all_files(monkeypatch, make_docx):
    monkeypatch.setattr(api_v2, "WORKERS", 1)  # window = 2
    data = make_docx(["Python"])
    r = post([("files", (f"{i}.docx", data)) for i in range(25)])
    assert len(lines(r)) == 25


# --- one real pool, to prove the pebble integration works end to end ---
def test_real_pebble_pool(make_docx):
    pool = ProcessPool(max_workers=1)
    api_v2.app.state.pool = pool
    try:
        r = post([("files", ("a.docx", make_docx(["Python and SQL"])))])
        (res,) = lines(r)
        assert res["status"] == "success" and res["score"] == 2
    finally:
        pool.stop()
        pool.join()
