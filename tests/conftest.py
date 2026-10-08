import io
import os
import sys

import docx
import pytest
from reportlab.pdfgen import canvas

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


@pytest.fixture
def make_pdf():
    def _make(text):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        y = 800
        for line in text.split("\n"):
            c.drawString(50, y, line)
            y -= 20
        c.save()
        return buf.getvalue()
    return _make


@pytest.fixture
def make_docx():
    def _make(paragraphs, table_cells=()):
        d = docx.Document()
        for p in paragraphs:
            d.add_paragraph(p)
        if table_cells:
            t = d.add_table(rows=1, cols=len(table_cells))
            for cell, txt in zip(t.rows[0].cells, table_cells):
                cell.text = txt
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()
    return _make
