"""Resume parser API, v2.

Same endpoint and response format as api_v1.py, plus:
  * per-file timeout: a stuck file is killed and reported, the worker is replaced
  * a crashed worker only fails the file it was running (no pool rebuild / leaked executors)
  * work is cancelled when the client disconnects
  * files are read lazily with a bounded in-flight window (bounded memory)
  * limit on concurrent requests (429) and on total request size (413)
  * optional API key (set API_KEY; clients send the X-API-Key header)
  * content sniffing (file bytes must match the extension), skill count/length limits
  * routes under /v2 (the unprefixed paths still work)

Configuration (environment variables): HOST, PORT, WORKERS, FILE_TIMEOUT_SEC,
MAX_CONCURRENT_REQUESTS, MAX_REQUEST_MB, API_KEY.
"""
import asyncio
import hmac
import io
import json
import logging
import os
import re
import time
from concurrent.futures import TimeoutError as FutureTimeout
from contextlib import asynccontextmanager
from functools import lru_cache
from typing import List, Optional

import docx
import pdfplumber
import pytesseract
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from pdf2image import convert_from_bytes
from pebble import ProcessExpired, ProcessPool
from PIL import Image

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("resume_parser")

# --- Configuration ---
WORKERS = int(os.getenv("WORKERS", max(1, (os.cpu_count() or 2) // 2)))
FILE_TIMEOUT_SEC = float(os.getenv("FILE_TIMEOUT_SEC", "60"))
MAX_CONCURRENT_REQUESTS = int(os.getenv("MAX_CONCURRENT_REQUESTS", "2"))
MAX_REQUEST_BYTES = int(float(os.getenv("MAX_REQUEST_MB", "256")) * 1024 * 1024)

# --- Limits ---
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB per file
MAX_RESUME_COUNT = 300
MAX_OCR_PAGES = 10
MAX_SKILLS = 50
MAX_SKILL_LENGTH = 64
SUPPORTED_EXTENSIONS = {"pdf", "docx", "jpg", "jpeg", "png"}
MAGIC = {
    "pdf": (b"%PDF",),
    "docx": (b"PK\x03\x04",),
    "png": (b"\x89PNG",),
    "jpg": (b"\xff\xd8",),
    "jpeg": (b"\xff\xd8",),
}

_active_requests = 0


# --- Worker side (runs in a child process) ---
@lru_cache(maxsize=32)
def build_skill_pattern(skills: tuple):
    """One regex for all skills. Handles multi-word skills and symbols like c++, c#, node.js."""
    parts = []
    for skill in sorted(set(skills), key=len, reverse=True):
        parts.append(r"\s+".join(re.escape(w) for w in skill.split()))
    return re.compile(r"(?<![\w+#])(?:" + "|".join(parts) + r")(?![\w+#])", re.IGNORECASE)


def extract_text(file_bytes, ext):
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            page_texts = [p.extract_text() or "" for p in pdf.pages]
        text = "\n".join(t for t in page_texts if t)
        if len(text.strip()) < 50:  # likely a scanned PDF -> OCR
            pages = convert_from_bytes(file_bytes, last_page=MAX_OCR_PAGES)
            text = "\n".join(pytesseract.image_to_string(p) for p in pages)
        return text
    if ext == "docx":
        d = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        for section in d.sections:  # headers/footers often hold contact info and skills
            parts.extend(p.text for p in section.header.paragraphs)
            parts.extend(p.text for p in section.footer.paragraphs)
        return "\n".join(parts)
    return pytesseract.image_to_string(Image.open(io.BytesIO(file_bytes)))


def process_single_resume(file_bytes, filename, target_skills):
    start_time = time.time()
    ext = os.path.splitext(filename)[1].lower().lstrip(".")

    if ext not in SUPPORTED_EXTENSIONS:
        return {"status": "error", "filename": filename, "error": f"Unsupported file type: '{ext}'"}
    if not file_bytes.startswith(MAGIC[ext]):
        return {"status": "error", "filename": filename, "error": "File content does not match its extension."}

    try:
        text = extract_text(file_bytes, ext)
        pattern = build_skill_pattern(tuple(target_skills))
        found = {" ".join(m.group(0).lower().split()) for m in pattern.finditer(text)}
        return {
            "status": "success",
            "filename": filename,
            "score": len(found),
            "matched_skills": sorted(found),
            "time_taken_sec": round(time.time() - start_time, 3),
        }
    except Exception:
        logger.exception("Failed to process %s", filename)
        return {"status": "error", "filename": filename, "error": "Could not process this file."}


# --- Server side ---
@asynccontextmanager
async def lifespan(app: FastAPI):
    # max_tasks recycles workers periodically so memory from big PDFs can't pile up
    app.state.pool = ProcessPool(max_workers=WORKERS, max_tasks=100)
    yield
    app.state.pool.stop()
    app.state.pool.join()


app = FastAPI(title="Resume Parser API v2", lifespan=lifespan)


@app.middleware("http")
async def limit_request_size(request: Request, call_next):
    length = request.headers.get("content-length")
    if length and length.isdigit() and int(length) > MAX_REQUEST_BYTES:
        return JSONResponse({"detail": "Request too large."}, status_code=413)
    return await call_next(request)


def require_api_key(x_api_key: Optional[str] = Header(None)):
    expected = os.getenv("API_KEY")
    if expected and not hmac.compare_digest((x_api_key or "").encode(), expected.encode()):
        raise HTTPException(status_code=401, detail="Invalid or missing API key.")


def error_result(name, message):
    return {"status": "error", "filename": name, "error": message}


async def run_one(pool, content, name, skills):
    """Run one file in the pool. Never raises (except cancellation): every failure becomes a result."""
    fut = pool.schedule(process_single_resume, args=(content, name, skills), timeout=FILE_TIMEOUT_SEC)
    try:
        return await asyncio.wrap_future(fut)
    except asyncio.CancelledError:
        fut.cancel()  # client went away: free (and, if running, kill) the worker
        raise
    except (FutureTimeout, TimeoutError):
        logger.warning("Timed out after %ss: %s", FILE_TIMEOUT_SEC, name)
        return error_result(name, f"Processing timed out after {FILE_TIMEOUT_SEC:g}s.")
    except ProcessExpired:
        logger.error("Worker died while processing %s", name)
        return error_result(name, "Worker crashed while processing this file.")
    except Exception:
        logger.exception("Unexpected failure for %s", name)
        return error_result(name, "Internal error.")


def parse_skills(raw: str):
    skills = list(dict.fromkeys(s.strip().lower() for s in raw.split(",") if s.strip()))
    if not skills:
        raise HTTPException(status_code=400, detail="No skills provided")
    if len(skills) > MAX_SKILLS:
        raise HTTPException(status_code=400, detail=f"Too many skills (max {MAX_SKILLS}).")
    if any(len(s) > MAX_SKILL_LENGTH for s in skills):
        raise HTTPException(status_code=400, detail=f"A skill is longer than {MAX_SKILL_LENGTH} characters.")
    return skills


@app.get("/health")
@app.get("/v2/health")
async def health():
    return {"status": "ok", "version": 2, "workers": WORKERS}


@app.post("/rank-resumes", dependencies=[Depends(require_api_key)])
@app.post("/v2/rank-resumes", dependencies=[Depends(require_api_key)])
async def rank_resumes(request: Request, skills: str = Form(...), files: List[UploadFile] = File(...)):
    global _active_requests

    if len(files) > MAX_RESUME_COUNT:
        raise HTTPException(
            status_code=400,
            detail=f"Too many files! Maximum allowed is {MAX_RESUME_COUNT}, but you sent {len(files)}.",
        )
    target_skills = parse_skills(skills)
    for f in files:
        if f.size is not None and f.size > MAX_FILE_SIZE:
            raise HTTPException(status_code=413, detail=f"File {f.filename} exceeds 5MB limit.")

    if _active_requests >= MAX_CONCURRENT_REQUESTS:
        raise HTTPException(status_code=429, detail="Server busy, try again shortly.",
                            headers={"Retry-After": "5"})

    pool = request.app.state.pool
    window = max(2, WORKERS * 2)  # max files read-but-unfinished at once
    _active_requests += 1

    async def stream_results():
        global _active_requests
        pending = set()
        try:
            async def drain(limit):
                nonlocal pending
                while len(pending) > limit:
                    done, pending = await asyncio.wait(pending, return_when=asyncio.FIRST_COMPLETED)
                    for d in done:
                        yield json.dumps(d.result()) + "\n"

            for f in files:
                name = f.filename or "unnamed"
                content = await f.read()
                if len(content) > MAX_FILE_SIZE:  # don't trust the client-declared size
                    yield json.dumps(error_result(name, "File exceeds 5MB limit.")) + "\n"
                    continue
                pending.add(asyncio.ensure_future(run_one(pool, content, name, target_skills)))
                async for line in drain(window - 1):
                    yield line
            async for line in drain(0):
                yield line
        finally:
            for t in pending:  # reached on client disconnect
                t.cancel()
            _active_requests -= 1

    return StreamingResponse(stream_results(), media_type="application/x-ndjson")


if __name__ == "__main__":
    import uvicorn
    # Localhost by default; set HOST=0.0.0.0 to expose on the network (and set API_KEY first).
    uvicorn.run(app, host=os.getenv("HOST", "127.0.0.1"), port=int(os.getenv("PORT", "8000")))
