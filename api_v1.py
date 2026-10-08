import io
import os
import re
import time
import json
import asyncio
import logging
from concurrent.futures import ProcessPoolExecutor
from concurrent.futures.process import BrokenProcessPool
from typing import List

import pdfplumber
import docx
import pytesseract
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from pdf2image import convert_from_bytes
from PIL import Image

logger = logging.getLogger("resume_parser")
app = FastAPI(title="Safe Real-Time Resume Parser")

# --- 1. WORKER LIMIT ---
half_cpu = int(os.getenv("WORKERS", max(1, (os.cpu_count() or 2) // 2)))
executor = ProcessPoolExecutor(max_workers=half_cpu)

# --- 2. GLOBAL LIMITS ---
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB per file
MAX_RESUME_COUNT = 300           # Strict limit on total files per request
MAX_OCR_PAGES = 10               # OCR at most this many pages per PDF
SUPPORTED_EXTENSIONS = {"pdf", "docx", "jpg", "jpeg", "png"}


def build_skill_pattern(target_skills):
    """One regex for all skills. Handles multi-word skills and symbols like c++, c#, node.js."""
    parts = []
    for skill in sorted(set(target_skills), key=len, reverse=True):
        words = skill.split()
        parts.append(r"\s+".join(re.escape(w) for w in words))
    return re.compile(r"(?<![\w+#])(?:" + "|".join(parts) + r")(?![\w+#])", re.IGNORECASE)


def extract_text(file_bytes, ext):
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            page_texts = [p.extract_text() or "" for p in pdf.pages]
        text = "\n".join(t for t in page_texts if t)
        if len(text.strip()) < 50:  # likely a scanned PDF -> OCR
            pages = convert_from_bytes(file_bytes, last_page=MAX_OCR_PAGES)
            text = "\n".join(pytesseract.image_to_string(p) for p in pages)
        return text
    if ext == "docx":
        doc = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    return pytesseract.image_to_string(Image.open(io.BytesIO(file_bytes)))


def process_single_resume(file_bytes, filename, target_skills):
    start_time = time.time()
    ext = os.path.splitext(filename)[1].lower().lstrip(".")

    if ext not in SUPPORTED_EXTENSIONS:
        return {"status": "error", "filename": filename, "error": f"Unsupported file type: '{ext}'"}

    try:
        text = extract_text(file_bytes, ext)
        pattern = build_skill_pattern(target_skills)
        found = {" ".join(m.group(0).lower().split()) for m in pattern.finditer(text)}

        return {
            "status": "success",
            "filename": filename,
            "score": len(found),
            "matched_skills": sorted(found),
            "time_taken_sec": round(time.time() - start_time, 3),
        }
    except Exception:
        logger.exception("Failed to process %s", filename)
        return {"status": "error", "filename": filename, "error": "Could not process this file."}


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.post("/rank-resumes")
async def rank_resumes(
    skills: str = Form(...),
    files: List[UploadFile] = File(...),
):
    if len(files) > MAX_RESUME_COUNT:
        raise HTTPException(
            status_code=400,
            detail=f"Too many files! Maximum allowed is {MAX_RESUME_COUNT}, but you sent {len(files)}.",
        )

    target_skills = [s.strip().lower() for s in skills.split(",") if s.strip()]
    if not target_skills:
        raise HTTPException(status_code=400, detail="No skills provided")

    for file in files:
        if file.size is not None and file.size > MAX_FILE_SIZE:
            raise HTTPException(status_code=413, detail=f"File {file.filename} exceeds 5MB limit.")

    async def run_one(loop, file_bytes, name):
        global executor
        try:
            return await loop.run_in_executor(executor, process_single_resume, file_bytes, name, target_skills)
        except BrokenProcessPool:
            logger.error("Worker pool broke while processing %s; rebuilding", name)
            executor = ProcessPoolExecutor(max_workers=half_cpu)
            return {"status": "error", "filename": name, "error": "Worker crashed while processing this file."}
        except Exception:
            logger.exception("Unexpected failure for %s", name)
            return {"status": "error", "filename": name, "error": "Internal error."}

    async def stream_results():
        loop = asyncio.get_running_loop()
        tasks = []
        for f in files:
            content = await f.read()
            if len(content) > MAX_FILE_SIZE:  # don't trust the client-declared size
                tasks.append(asyncio.ensure_future(asyncio.sleep(0, {
                    "status": "error", "filename": f.filename, "error": "File exceeds 5MB limit."})))
                continue
            tasks.append(asyncio.ensure_future(run_one(loop, content, f.filename)))

        for task in asyncio.as_completed(tasks):
            yield json.dumps(await task) + "\n"

    return StreamingResponse(stream_results(), media_type="application/x-ndjson")


if __name__ == "__main__":
    import uvicorn
    # Localhost by default; set HOST=0.0.0.0 to expose on the network (add auth first).
    uvicorn.run(app, host=os.getenv("HOST", "127.0.0.1"), port=int(os.getenv("PORT", "8000")))
